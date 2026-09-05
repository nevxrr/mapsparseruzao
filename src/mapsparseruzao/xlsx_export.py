from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from mapsparseruzao.constants import XLSX_HEADERS
from mapsparseruzao.models import Company


def export_xlsx(companies: Sequence[Company], path: str | Path) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ЮЗАО без сайта"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    sheet.append(list(XLSX_HEADERS))
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for company in companies:
        sheet.append(list(company.to_row()))

    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"

    widths = (36, 14, 18, 28, 46, 22, 28, 32, 12, 28, 12, 12, 42, 16)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(destination)
    return destination

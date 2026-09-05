from pathlib import Path

from openpyxl import load_workbook

from mapsparseruzao.cli import run
from mapsparseruzao.constants import XLSX_HEADERS
from mapsparseruzao.models import Company
from mapsparseruzao.xlsx_export import export_xlsx


def test_export_xlsx_writes_headers_and_rows(tmp_path: Path) -> None:
    company = Company(
        name="Ателье",
        category="craft",
        kind="tailor",
        district="Коньково",
        address="Москва, ул. Профсоюзная, 100",
        phone="+7 495 111-11-11",
        email="",
        website="",
        opening_hours="Mo-Fr 10:00-19:00",
        lat=55.64,
        lon=37.52,
        osm_url="https://www.openstreetmap.org/node/123",
    )
    path = export_xlsx([company], tmp_path / "out.xlsx")
    workbook = load_workbook(path)
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == list(XLSX_HEADERS)
    assert sheet["A2"].value == "Ателье"
    assert sheet["I2"].value == "нет"
    assert sheet["M2"].value == company.osm_url


def test_cli_saves_filtered_xlsx(tmp_path: Path, monkeypatch) -> None:
    payload = {
        "elements": [
            {
                "type": "node",
                "id": 10,
                "lat": 55.65,
                "lon": 37.55,
                "tags": {
                    "name": "Салон без сайта",
                    "shop": "beauty",
                    "addr:suburb": "Ясенево",
                    "phone": "+7 495 222-22-22",
                },
            },
            {
                "type": "node",
                "id": 11,
                "lat": 55.66,
                "lon": 37.56,
                "tags": {"name": "Салон с сайтом", "shop": "beauty", "website": "https://yes.ru"},
            },
        ]
    }

    class _Client:
        def __init__(self, *args, **kwargs) -> None:
            pass

        def fetch_yuzao(self, **kwargs):
            return payload

    monkeypatch.setattr("mapsparseruzao.cli.OverpassClient", _Client)
    output = tmp_path / "yuzao.xlsx"
    assert run(["--output", str(output), "--categories", "shop"]) == 0

    workbook = load_workbook(output)
    names = [row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)]
    assert names == ["Салон без сайта"]
